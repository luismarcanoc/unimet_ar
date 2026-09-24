from __future__ import annotations

import csv
from pathlib import Path

from .models import Clase


PUBLIC_COLUMNS = [
    "DIAS",
    "HORA",
    "AULA",
    "MATERIA(ASIGNATURA)",
    "SECCION",
    "PROFESOR",
]
AUDIT_COLUMNS = PUBLIC_COLUMNS + ["CODIGO", "PERIODO", "ARCHIVO"]


def guardar_csv(path: Path, clases: list[Clase], auditoria: bool = False) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    columns = AUDIT_COLUMNS if auditoria else PUBLIC_COLUMNS
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=columns)
        writer.writeheader()
        for clase in clases:
            writer.writerow(clase.fila_auditoria() if auditoria else clase.fila_publica())


def guardar_excel(path: Path, grupos: dict[str, list[Clase]]) -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError as exc:
        raise RuntimeError(
            "Falta openpyxl. Ejecuta: python -m pip install -r requirements.txt"
        ) from exc

    workbook = Workbook()
    workbook.remove(workbook.active)
    for name, clases in grupos.items():
        sheet = workbook.create_sheet(title=name[:31])
        sheet.append(PUBLIC_COLUMNS)
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="C44720")
            cell.alignment = Alignment(horizontal="center")
        for clase in clases:
            row = clase.fila_publica()
            sheet.append([row[column] for column in PUBLIC_COLUMNS])
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for index, column in enumerate(PUBLIC_COLUMNS, start=1):
            max_length = max(len(str(cell.value or "")) for cell in sheet[get_column_letter(index)])
            sheet.column_dimensions[get_column_letter(index)].width = min(max_length + 2, 55)
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)

